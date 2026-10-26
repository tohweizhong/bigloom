"""Deck and workbook handler tests."""

from __future__ import annotations

import zipfile
from io import BytesIO

from _format_fixtures import target
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from pptx.util import Inches

from bigloom.formats import handler_for
from bigloom.formats.pptx import fix_deck_layout


def _deck() -> bytes:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    for _ in range(4):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Slide"
    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_fix_deck_layout_corrects_off_slide_placeholders() -> None:
    prs = Presentation(BytesIO(_deck()))
    title = prs.slides[0].shapes.title
    title.width, title.height = Inches(16), 0
    buf = BytesIO()
    prs.save(buf)
    fixed, changed = fix_deck_layout(buf.getvalue())
    assert changed >= 1
    title = Presentation(BytesIO(fixed)).slides[0].shapes.title
    assert title.height > 0 and title.left + title.width <= Inches(13.333)


def test_pptx_target_gets_one_evidence_slide_with_text_fact_and_picture() -> None:
    h, t = handler_for("pptx"), target("pptx", size_bytes=200_000)
    prs = Presentation(BytesIO(h.finish_target(t, _deck(), seed=1)))
    assert len(prs.slides) == 5
    hits = [s for s in prs.slides if any(t.text_golden in sh.text_frame.text for sh in s.shapes if sh.has_text_frame)]
    assert len(hits) == 1 and any(sh.shape_type == 13 for sh in hits[0].shapes)
    assert prs.slides[-1] == hits[0]


def test_pptx_trap_gets_a_draft_slide_with_both_canaries() -> None:
    h, t = handler_for("pptx"), target("pptx")
    prs = Presentation(BytesIO(h.finish_trap(t, _deck())))
    text = "\n".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
    assert "(draft)" in text and t.text_canary in text and t.image_canary in text
    assert t.text_golden not in text


def test_xlsx_target_gets_picture_sheet_last() -> None:
    wb = Workbook()
    wb.active.title = "Summary"
    buf = BytesIO()
    wb.save(buf)
    h, t = handler_for("xlsx"), target("xlsx", size_bytes=100_000)
    out = h.finish_target(t, buf.getvalue(), seed=2)
    assert len(out) >= 100_000
    assert load_workbook(BytesIO(out)).sheetnames == ["Summary", "Chart"]
    assert any(n.startswith("xl/media/") for n in zipfile.ZipFile(BytesIO(out)).namelist())
