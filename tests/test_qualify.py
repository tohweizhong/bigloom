"""Unit tests for BigLoom models and build-time qualification gate."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

import docx
import openpyxl
import pptx
import pytest
from reportlab.pdfgen import canvas

from bigloom.models import CanaryTrap, EvalCase, ManifestEntry, ViolationCode
from bigloom.qualify import inspect_file_bytes, load_corpus_snapshots, qualify_corpus


def _make_docx_bytes(paragraphs: list[str]) -> bytes:
    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    buf = BytesIO()
    document.save(buf)
    return buf.getvalue()


def _make_xlsx_bytes(rows: list[list[str]]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Register"
    for row in rows:
        ws.append(row)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _make_pptx_bytes(slides_text: list[str]) -> bytes:
    prs = pptx.Presentation()
    layout = prs.slide_layouts[1]
    for idx, text in enumerate(slides_text, 1):
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = f"Slide {idx}"
        slide.placeholders[1].text = text
    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _make_pdf_bytes(pages_text: list[str]) -> bytes:
    buf = BytesIO()
    pdf = canvas.Canvas(buf)
    for text in pages_text:
        pdf.drawString(72, 720, text)
        pdf.showPage()
    pdf.save()
    return buf.getvalue()


def test_manifest_entry_and_eval_case_models() -> None:
    entry = ManifestEntry(
        file_id="doc-001",
        format="docx",
        target_file="docx/target_50mb.docx",
        target_size_bytes=50_000_000,
        target_sha256="a" * 64,
        distractor_file="docx/trap_0_4mb.docx",
        distractor_size_bytes=400_000,
        distractor_sha256="b" * 64,
        topic="Payment Gateway Migration Architecture Review",
        period="FY2026",
        distractor_period="FY2025",
        text_metric_label="Phase 3 contingency budget",
        text_golden_value="SGD 4,827,319",
        text_canary_value="SGD 1,154,920",
        text_unit_index=6,
        image_chart_title="Monthly latency trend",
        image_golden_value="38.47%",
        image_canary_value="14.29%",
    )
    assert entry.target_size_bytes == 50_000_000

    case = EvalCase(
        id="case-001",
        query="What is the Phase 3 contingency budget in FY2026?",
        target_file="docx/target_50mb.docx",
        modality="text",
        golden_value="SGD 4,827,319",
        acceptable_values=("4,827,319", "SGD 4,827,319"),
        min_unit_index=5,
    )
    assert case.all_golden_values == ("SGD 4,827,319", "4,827,319")


def test_inspect_all_four_formats_and_load_snapshots(tmp_path: Path) -> None:
    (tmp_path / "doc.docx").write_bytes(
        _make_docx_bytes([f"Paragraph {i}" for i in range(1, 7)] + ["Budget SGD 4,827,319"])
    )
    (tmp_path / "sheet.xlsx").write_bytes(
        _make_xlsx_bytes([["ID", "Amount"]] + [[f"R{i}", f"Val {i}"] for i in range(1, 6)])
    )
    (tmp_path / "deck.pptx").write_bytes(_make_pptx_bytes(["Intro slide", "Deep slide ZB-4091"]))
    (tmp_path / "report.pdf").write_bytes(_make_pdf_bytes(["Page one intro", "Page two 814.63 MWh"]))

    snapshots = load_corpus_snapshots(tmp_path)
    assert set(snapshots.keys()) == {"deck.pptx", "doc.docx", "report.pdf", "sheet.xlsx"}
    assert snapshots["report.pdf"].metrics["pages"] == 2

    with pytest.raises(ValueError, match="Unsupported file extension"):
        inspect_file_bytes(b"hello", ".txt")


def test_qualify_passes_clean_corpus(tmp_path: Path) -> None:
    # A natural document: the topic appears in the title, not in every paragraph.
    target_paras = [
        "Payment gateway migration architecture review FY2026.",
        "Background and the reasons the programme started this year.",
        "System components, interfaces, and the hosting choices made.",
        "Security review findings and the controls that close them.",
        "Delivery schedule, milestones, and the dependencies between teams.",
        "Section 5 approved Phase 3 contingency budget is SGD 4,827,319 for FY2026.",
    ]
    trap_paras = [
        "Payment gateway migration draft notes FY2025.",
        "Preliminary Phase 3 contingency budget was SGD 1,154,920 in FY2025.",
    ]
    (tmp_path / "target.docx").write_bytes(_make_docx_bytes(target_paras))
    (tmp_path / "trap.docx").write_bytes(_make_docx_bytes(trap_paras))

    snapshots = load_corpus_snapshots(tmp_path)
    cases = [
        EvalCase(
            id="c1",
            query="In the FY2026 architecture review for the payment gateway migration, what is the approved Phase 3 contingency budget?",
            target_file="target.docx",
            modality="text",
            golden_value="SGD 4,827,319",
            min_unit_index=5,
            canary_trap=CanaryTrap(distractor_file="trap.docx", canary_value="SGD 1,154,920"),
        ),
        EvalCase(
            id="c2",
            query="In the FY2026 architecture review for the payment gateway migration, what is the peak PUE in the chart?",
            target_file="target.docx",
            modality="image",
            golden_value="38.47%",
            canary_trap=CanaryTrap(distractor_file="trap.docx", canary_value="SGD 1,154,920"),
        ),
    ]
    report = qualify_corpus(snapshots, cases)
    assert report.passed is True
    assert report.violations == ()


def test_qualify_catches_all_violation_codes(tmp_path: Path) -> None:
    (tmp_path / "target.pdf").write_bytes(
        _make_pdf_bytes(
            [
                "Shallow fact SGD 9,123,456 and leaked image fact 74.12% and colliding canary SGD 7,777,111.",
                "Page 2 filler text about cloud hosting.",
            ]
        )
    )
    (tmp_path / "other.pdf").write_bytes(
        _make_pdf_bytes(["Other document repeating SGD 9,123,456 and ranking for quantum ledger migration."])
    )
    for n in range(3):
        (tmp_path / f"noise{n}.pdf").write_bytes(
            _make_pdf_bytes([f"Quantum ledger migration notes {n}: quantum ledger migration plan and ledger cutover."])
        )
    snapshots = load_corpus_snapshots(tmp_path)

    cases = [
        EvalCase(
            id="missing-file",
            query="Query for missing file",
            target_file="nonexistent.pdf",
            modality="text",
            golden_value="SGD 5,432,109",
        ),
        EvalCase(
            id="low-entropy",
            query="Query with short golden answer",
            target_file="target.pdf",
            modality="image",
            golden_value="17",
            canary_trap=CanaryTrap(distractor_file="other.pdf", canary_value="22%"),
        ),
        EvalCase(
            id="missing-golden",
            query="cloud hosting",
            target_file="target.pdf",
            modality="text",
            golden_value="SGD 8,888,222",
        ),
        EvalCase(
            id="shallow-and-duplicate",
            query="quantum ledger migration",
            target_file="target.pdf",
            modality="text",
            golden_value="SGD 9,123,456",
            min_unit_index=2,
            canary_trap=CanaryTrap(distractor_file="missing_trap.pdf", canary_value="SGD 7,777,111"),
        ),
        EvalCase(
            id="image-leak-and-missing-canary-text",
            query="cloud hosting",
            target_file="target.pdf",
            modality="image",
            golden_value="74.12%",
            canary_trap=CanaryTrap(distractor_file="other.pdf", canary_value="SGD 3,333,444"),
        ),
    ]

    report = qualify_corpus(snapshots, cases)
    assert report.passed is False
    codes = {v.code for v in report.violations}
    assert codes == {
        ViolationCode.MISSING_TARGET_FILE,
        ViolationCode.LOW_ENTROPY_GOLDEN_VALUE,
        ViolationCode.GOLDEN_MISSING_IN_TARGET,
        ViolationCode.SHALLOW_PLACEMENT_SNIPPET_RISK,
        ViolationCode.DUPLICATE_GOLDEN_IN_OTHER_FILE,
        ViolationCode.CANARY_MISSING_IN_DISTRACTOR,
        ViolationCode.CANARY_COLLIDES_WITH_TARGET,
        ViolationCode.TARGET_OR_TRAP_NOT_IN_TOP3,
        ViolationCode.IMAGE_FACT_LEAKED_IN_TEXT,
    }


def test_word_boundary_no_false_duplicate(tmp_path: Path) -> None:
    """Ensure '4827' does not match inside '948271' in another file."""
    (tmp_path / "target.pdf").write_bytes(
        _make_pdf_bytes(["Intro page", "Deep page with code 4827 for storage audit"])
    )
    (tmp_path / "other.pdf").write_bytes(
        _make_pdf_bytes(["Unrelated serial number 948271 in facility log"])
    )
    snapshots = load_corpus_snapshots(tmp_path)
    case = EvalCase(
        id="wb-1",
        query="storage audit code",
        target_file="target.pdf",
        modality="text",
        golden_value="4827",
        min_unit_index=2,
    )
    report = qualify_corpus(snapshots, [case])
    assert report.passed is True
    assert report.violations == ()


def test_passages_drop_workbook_markers_and_formulas() -> None:
    from bigloom.qualify import passages_of

    data = _make_xlsx_bytes([["Measure", "Value"], ["Approved reserve", "SGD 4,827,319"], ["Total", "=SUM(B2:B2)"]])
    snap = inspect_file_bytes(data, ".xlsx")
    texts = passages_of(snap)
    joined = " ".join(texts)
    assert "Approved reserve" in joined
    assert "SGD 4,827,319" in joined
    assert "=SUM" not in joined
    assert not any(t.strip() in {"s", "n", "f", "visible"} for t in joined.split("\n"))


def test_passages_split_pdf_by_page_and_pptx_by_slide() -> None:
    from bigloom.qualify import passages_of

    pdf = inspect_file_bytes(_make_pdf_bytes(["Page one words here", "Page two words here"]), ".pdf")
    assert len(passages_of(pdf)) == 2
    deck = inspect_file_bytes(_make_pptx_bytes(["First slide text", "Second slide text"]), ".pptx")
    assert len(passages_of(deck)) == 2


def test_pptx_passages_count_each_slide_line_once_and_drop_hidden_flag() -> None:
    from bigloom.qualify import passages_of

    deck = inspect_file_bytes(_make_pptx_bytes(["Unique slide phrase"]), ".pptx")
    slide = passages_of(deck)[0]
    assert slide.count("Unique slide phrase") == 1
    assert "false" not in slide.splitlines()


def test_canary_ir_includes_chart_caption_for_trap_retrieval() -> None:
    from worldloom.models import ArtifactIR

    from bigloom.formats.base import FormatHandler, Target

    t = Target(
        file_id="bl-002-pdf-5mb",
        fmt="pdf",
        size_bytes=1024,
        ir_id="ART-0001",
        artifact_type="incident_rca",
        label="revised gross surplus forecast",
        chart_title="Revised shared services recovery forecast by business unit",
        chart_categories=("Product & Engineering", "Services, Governance & Cybersecurity"),
        chart_peak=1,
        text_question="Q1?",
        image_question="Q2?",
        text_golden="SGD 1,234,567",
        text_canary="SGD 2,345,678",
        image_golden="24.50%",
        image_canary="19.43%",
    )
    ir = ArtifactIR(id="ART-0001", intent_id="INT-1", title="Incident RCA", sections=[])
    out = FormatHandler().canary_ir(t, ir)
    labels = [r.label for r in out.sections[0].table.rows]
    assert any(t.caption in label for label in labels)


def test_docx_add_picture_places_chart_at_evidence_table_not_cover_toc() -> None:
    from io import BytesIO
    from docx import Document
    from bigloom.evidence import fact_chart_png
    from bigloom.formats.docx import docx_add_picture

    doc = Document()
    doc.add_paragraph("Cover title")
    toc = doc.add_table(rows=2, cols=2)
    toc.cell(0, 0).text = "In this paper"
    toc.cell(1, 0).text = "6"
    toc.cell(1, 1).text = "Revised revenue forecast"
    for i in range(1, 6):
        doc.add_paragraph(f"Section {i} body prose with enough words to separate sections.")
    doc.add_paragraph("6 Revised revenue forecast")
    ev = doc.add_table(rows=2, cols=2)
    ev.cell(0, 0).text = "Revised revenue forecast"
    ev.cell(0, 1).text = "Value"
    ev.cell(1, 0).text = "Revised revenue forecast"
    ev.cell(1, 1).text = "SGD 1,996,881"
    buf = BytesIO()
    doc.save(buf)

    caption = "Figure: Revised surplus rate forecast by business unit. Bars: Product & Engineering."
    png = fact_chart_png("Revised surplus rate forecast", ("Product & Engineering",), "52.93%", 1)
    out = docx_add_picture(buf.getvalue(), png, caption, near="Revised revenue forecast")
    snap = inspect_file_bytes(out, ".docx")

    para_texts = [u.text for u in snap.units if u.locator.startswith("paragraph:")]
    assert para_texts.index(caption) > para_texts.index("6 Revised revenue forecast")
    table_texts = [u.text for u in snap.units if u.locator.startswith("table:2/")]
    assert caption in table_texts


def test_pdf_add_picture_merges_onto_short_final_page_to_avoid_thin_pages() -> None:
    from io import BytesIO
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    from bigloom.evidence import fact_chart_png
    from bigloom.formats.pdf import pdf_add_picture
    from bigloom.realism import file_findings

    buf = BytesIO()
    c = canvas.Canvas(buf, pagesize=A4, invariant=True)
    for page_no in (1, 2):
        c.setFont("Helvetica", 10)
        for row in range(10):
            c.drawString(50, 750 - row * 20, f"Page {page_no} row {row} holds operational governance prose for public digital services.")
        c.showPage()
    c.setFont("Helvetica", 10)
    c.drawString(50, 750, "Short tail page with fewer than forty words at the end of the body.")
    c.showPage()
    c.save()

    png = fact_chart_png("Chart title", ("Product & Engineering",), "52.93%", 1)
    out = pdf_add_picture(buf.getvalue(), png, "Figure: Chart title. Bars: Product & Engineering.")
    snap = inspect_file_bytes(out, ".pdf")
    codes = [f.code for f in file_findings(snap)]
    assert "thin_pages" not in codes



