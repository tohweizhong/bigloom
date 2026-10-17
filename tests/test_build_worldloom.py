"""Phase 4 tests: WorldLoom renders each target, and BigLoom adds only the test evidence.

Track worldloom_native_bigloom_20261016, FR4.
"""

from __future__ import annotations

import json
import zipfile
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

from bigloom.build import build_corpus, style_seed_for
from bigloom.models import ManifestEntry
from bigloom.qualify import load_corpus_snapshots

SIZE_MB = 0.3


@pytest.fixture(scope="module")
def built(tmp_path_factory: pytest.TempPathFactory) -> tuple[Path, list[ManifestEntry]]:
    out = tmp_path_factory.mktemp("phase4") / "out"
    entries = build_corpus(out, sizes_mb=(SIZE_MB,), seed=42)
    return out, entries


def _entry(entries: list[ManifestEntry], fmt: str) -> ManifestEntry:
    return next(e for e in entries if e.format == fmt)


def test_each_target_is_a_different_worldloom_document(built: tuple[Path, list[ManifestEntry]]) -> None:
    _, entries = built
    assert [e.format for e in entries] == ["docx", "xlsx", "pptx", "pdf"]
    assert len({e.document_type for e in entries}) == 4
    assert all(e.document_title for e in entries)
    assert len({e.text_metric_label for e in entries}) == 4


def test_docx_has_cover_contents_sections_and_picture(built: tuple[Path, list[ManifestEntry]]) -> None:
    out, entries = built
    e = _entry(entries, "docx")
    doc = Document(str(out / e.target_file))
    styles = [(p.style.name, p.text.strip()) for p in doc.paragraphs]
    assert ("Title", e.document_title) in styles
    assert any(text == "Contents" for style, text in styles if style.startswith("Heading"))
    assert sum(1 for style, _ in styles if style == "Heading 1") >= 4
    assert len(doc.tables) >= 3
    assert len(doc.inline_shapes) >= 1


def test_xlsx_keeps_native_charts_and_adds_picture_sheet(built: tuple[Path, list[ManifestEntry]]) -> None:
    out, entries = built
    e = _entry(entries, "xlsx")
    data = (out / e.target_file).read_bytes()
    names = zipfile.ZipFile(BytesIO(data)).namelist()
    assert any(n.startswith("xl/charts/chart") for n in names)
    assert any(n.startswith("xl/media/") and n.endswith(".png") for n in names)
    sheets = load_workbook(BytesIO(data), read_only=True).sheetnames
    assert sheets[-1] == "Chart"
    assert (e.text_metric_label[:1].upper() + e.text_metric_label[1:])[:31] in sheets
    assert len(sheets) >= 8


def test_pptx_layout_is_fixed_and_evidence_slide_is_present(built: tuple[Path, list[ManifestEntry]]) -> None:
    out, entries = built
    e = _entry(entries, "pptx")
    prs = Presentation(str(out / e.target_file))
    width = prs.slide_width
    for container in [*prs.slide_layouts, *prs.slides]:
        for ph in container.placeholders:
            assert ph.height and ph.width, f"placeholder {ph.name} has no size"
            assert ph.left + ph.width <= width, f"placeholder {ph.name} runs off the slide"
    evidence = [s for s in prs.slides if any(e.text_golden_value in sh.text_frame.text for sh in s.shapes if sh.has_text_frame)]
    assert len(evidence) == 1
    assert any(sh.shape_type == 13 for sh in evidence[0].shapes)  # 13 = picture
    assert len(prs.slides) >= 10


def test_pdf_has_cover_and_picture_page(built: tuple[Path, list[ManifestEntry]]) -> None:
    out, entries = built
    e = _entry(entries, "pdf")
    reader = PdfReader(str(out / e.target_file))
    assert len(reader.pages) >= 8
    assert e.document_title in reader.pages[0].extract_text()
    assert any(page.images for page in reader.pages)


def test_text_fact_is_deep_and_image_fact_is_only_in_the_picture(built: tuple[Path, list[ManifestEntry]]) -> None:
    out, entries = built
    snaps = load_corpus_snapshots(out)
    for e in entries:
        assert e.text_unit_index >= 6, e.file_id
        text = "\n".join(u.text for u in snaps[e.target_file].units)
        assert e.text_golden_value in text.replace("\n", " ") or e.text_golden_value.replace(" ", "") in text.replace(" ", "")
        assert e.image_golden_value not in text, e.file_id
        assert e.text_canary_value not in text and e.image_canary_value not in text, e.file_id


def test_trap_is_the_draft_of_the_same_document_with_the_canary(built: tuple[Path, list[ManifestEntry]]) -> None:
    out, entries = built
    snaps = load_corpus_snapshots(out)
    for e in entries:
        trap = snaps[e.distractor_file]
        text = "\n".join(u.text for u in trap.units)
        flat = text.replace(" ", "").replace("\n", "")
        assert e.text_canary_value.replace(" ", "") in flat, e.file_id
        assert e.image_canary_value in flat, e.file_id
        assert e.text_golden_value.replace(" ", "") not in flat and e.image_golden_value not in flat
        assert e.document_title in text, e.file_id
        assert e.distractor_size_bytes < 1024 * 1024
        if e.format == "docx":
            assert "0.1" in text  # the WorldLoom v0.1 revision
        else:
            assert "(draft)" in text


def test_each_target_reaches_its_size(built: tuple[Path, list[ManifestEntry]]) -> None:
    out, entries = built
    target = int(SIZE_MB * 1024 * 1024)
    for e in entries:
        size = (out / e.target_file).stat().st_size
        assert size == e.target_size_bytes
        assert target <= size <= target + 4096, e.file_id


def test_documents_get_their_own_style_seed(built: tuple[Path, list[ManifestEntry]]) -> None:
    out, entries = built
    seeds = [e.style_seed for e in entries]
    assert None not in seeds and len(set(seeds)) == 4
    world_seed = entries[0].world_seed
    assert world_seed is not None
    # Each seed gives a different WorldLoom style genome, and the Word target uses its own.
    from worldloom.models import ArtifactIR
    from worldloom.render import docx as docx_render
    from worldloom.render import fonts

    def genome(seed: int | str):  # type: ignore[no-untyped-def]
        ir = ArtifactIR(id="ART-X", intent_id="INT-X", title="X", sections=[], metadata={"style_seed": str(seed)})
        return docx_render._genome_for(ir)

    keys = {genome(s).key for s in seeds if s is not None}
    assert len(keys) == 4
    assert genome(world_seed).key not in keys
    word = _entry(entries, "docx")
    assert word.style_seed is not None
    faces = fonts.named(genome(word.style_seed).typeface)
    doc = Document(str(out / word.target_file))
    assert doc.styles["Normal"].font.name == faces.body
    assert style_seed_for(world_seed, "ART-0001") != style_seed_for(world_seed, "ART-0002")


def test_manifest_holds_natural_questions(built: tuple[Path, list[ManifestEntry]]) -> None:
    out, entries = built
    cases = json.loads((out / "cases.json").read_text(encoding="utf-8"))
    queries = [c["query"] for c in cases]
    assert len(set(queries)) == 8
    for q in queries:
        assert "Final Audit Report" not in q
        assert q.endswith("?")
    for e in entries:
        assert e.document_title in e.text_question and e.document_title in e.image_question
        assert e.text_metric_label in e.text_question
        assert e.image_chart_title in e.image_question
