"""Synthesize large enterprise documents and paired canary trap files."""

from __future__ import annotations

import hashlib
import random
from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Literal

import xlsxwriter
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor as PptxRGBColor
from pptx.util import Inches as PptxInches
from pptx.util import Pt as PptxPt
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (
    Image as RLImage,
)
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from .charts import (
    CHART_KINDS,
    ChartContext,
    FigurePlate,
    encode_figure_tight,
    make_chart_figures,
)
from .corpus import CorpusProvenance, resolve_corpus
from .models import ManifestEntry
from .prose_writer import (
    MARKING,
    THEME_PALETTE,
    ProseCorpus,
    Writer,
    format_money,
    sheet_definition,
)
from .queries import generate_queries

SHEET_KINDS: tuple[str, ...] = ("budget", "procurement", "pipeline", "auth")
DATA_SHEET_PREFIX: dict[str, str] = {
    "budget": "Lines",
    "procurement": "POs",
    "pipeline": "Runs",
    "auth": "Txns",
}


@dataclass(frozen=True)
class _WorldContext:
    """Grounded enterprise entities and currency extracted from WorldLoom."""

    topics: tuple[str, ...]
    currency: str
    governance_line: str
    prose_corpus: ProseCorpus
    provenance: CorpusProvenance


def _resolve_world_context(
    *,
    seed: int,
    out_dir: Path,
    corpus_dir: Path | None = None,
    world_pack: str | None = None,
    domain: str | None = None,
    locale: str | None = None,
    facets: Mapping[str, str] | None = None,
    writer: str = "template",
) -> _WorldContext:
    """Load or build the WorldLoom corpus, then extract company, entities, and currency."""
    world, provenance = resolve_corpus(
        corpus_dir=corpus_dir,
        out_dir=out_dir,
        seed=seed,
        world_pack=world_pack,
        domain=domain,
        locale=locale,
        facets=facets,
        writer=writer,
    )

    prose_corpus = ProseCorpus.from_world(world)
    company_name = prose_corpus.company
    currency = prose_corpus.currency
    wl_topics: list[str] = [
        f"{company_name} {bu.name}" for bu in world.business_units
    ] + [
        f"{company_name} {label}" for label in prose_corpus.effective_part_labels()
    ]
    topics = list(wl_topics)

    cc_tags = [f"{cc.id} ({cc.name})" for cc in world.cost_centres[:3]]
    sys_tags = [f"{sys.id} ({sys.name})" for sys in world.systems[:3]]
    site_tags = [site.name for site in world.sites[:3]]
    fact_tags = [f"{fact.id} ({fact.kind})" for fact in world.facts[:3]]
    details: list[str] = [f"Company: {company_name}"]
    if cc_tags:
        details.append(f"Cost Centres: {', '.join(cc_tags)}")
    if sys_tags:
        details.append(f"Systems: {', '.join(sys_tags)}")
    if site_tags:
        details.append(f"Sites: {', '.join(site_tags)}")
    if fact_tags:
        details.append(f"Canonical Facts: {', '.join(fact_tags)}")
    governance_line = " | ".join(details) + "."

    return _WorldContext(
        topics=tuple(topics),
        currency=currency,
        governance_line=governance_line,
        prose_corpus=prose_corpus,
        provenance=provenance,
    )


def _split_image_budgets(target_bytes: int, max_chunk: int | None = None) -> list[int]:
    """Split a byte budget (up to 1 GiB) into raw RGB image sizes with an 8% safety margin."""
    if max_chunk is None:
        if target_bytes >= 250 * 1024 * 1024:
            max_chunk = 32_000_000
        elif target_bytes >= 50 * 1024 * 1024:
            max_chunk = 16_000_000
        else:
            max_chunk = 3_000_000
    padded = max(24_576, int(target_bytes * 1.08))
    chunks: list[int] = []
    remaining = padded
    while remaining > 0:
        step = min(remaining, max_chunk)
        chunks.append(step)
        remaining -= step
    return chunks


def _build_encoded_figures(
    rng: random.Random,
    *,
    corpus: ProseCorpus,
    header_line: str,
    chart_title: str,
    chart_value: str,
    target_bytes: int,
    theme: tuple[int, int, int],
) -> list[FigurePlate]:
    """Render rich high-resolution chart figures and encode each chunk to hit target_bytes."""
    chunks = _split_image_budgets(target_bytes)
    ctx = ChartContext(
        rng=rng,
        corpus=corpus,
        title=header_line,
        date="14 Oct 2026",
        theme=theme,
        chart_title=chart_title,
        printed_value=chart_value,
    )
    preferred = list(CHART_KINDS)
    rng.shuffle(preferred)
    figures = make_chart_figures(ctx, count=len(chunks), preferred_kinds=preferred)
    scales: dict[int, float] = {}
    for fig, chunk_bytes in zip(figures, chunks):
        encode_figure_tight(fig, target_bytes=chunk_bytes, scales=scales)
    return figures


def _document_sections(corpus: ProseCorpus) -> list[str]:
    """Return 7 document section headings from corpus.sections with Financial overview at slot 4."""
    others = [s for s in corpus.sections if s != "Financial overview"][:6]
    while len(others) < 6:
        others.append(f"Section {len(others) + 1}")
    return [*others[:4], "Financial overview", *others[4:6]]


def _plan_section_figures(
    sections: Sequence[str],
    figures: Sequence[FigurePlate],
) -> list[list[FigurePlate]]:
    """Assign figures across document sections while keeping Financial overview text-only."""
    slots = [i for i, s in enumerate(sections) if s != "Financial overview"]
    out: list[list[FigurePlate]] = [[] for _ in sections]
    if not slots:
        return out
    for idx, fig in enumerate(figures):
        out[slots[idx % len(slots)]].append(fig)
    return out


def _render_docx(
    rng: random.Random,
    *,
    corpus: ProseCorpus,
    header_line: str,
    governance_line: str,
    fact_sentence: str,
    chart_title: str,
    chart_value: str,
    target_bytes: int,
    theme: tuple[int, int, int] = (0, 84, 147),
) -> tuple[bytes, int]:
    """Render a multi-section .docx with headers, risk table, deep fact, and rich chart plates."""
    figures = _build_encoded_figures(
        rng,
        corpus=corpus,
        header_line=header_line,
        chart_title=chart_title,
        chart_value=chart_value,
        target_bytes=target_bytes,
        theme=theme,
    )
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    doc.sections[0].header.paragraphs[0].text = MARKING
    doc.sections[0].footer.paragraphs[0].text = f"{header_line[:60]}  |  {corpus.company}"

    doc.add_heading(header_line, 0)
    owner, role = rng.choice(corpus.people) if corpus.people else ("Tan Wei Jie", "Director")
    meta = doc.add_paragraph()
    meta.add_run(f"Prepared by {owner}, {role}. Date: 14 Oct 2026. {governance_line}").italic = True
    red = doc.add_paragraph().add_run(
        "This document is synthetic test data. All names and figures are invented."
    )
    red.font.color.rgb = RGBColor(0xAA, 0, 0)

    w = Writer(rng, corpus)
    sections = _document_sections(corpus)
    planned = _plan_section_figures(sections, figures)
    fig_no = 0
    text_unit_index = 6

    for sec_idx, (heading, sec_figs) in enumerate(zip(sections, planned)):
        doc.add_heading(heading, level=1)
        if heading == "Financial overview":
            doc.add_paragraph(
                "The list below gives the key figures for this document, "
                "and revenue from agency charges for selected lines of work."
            )
            doc.add_paragraph(fact_sentence, style="List Bullet")
            text_unit_index = len(doc.paragraphs)
            for (subj, per), (actual, budget) in sorted(corpus.revenue.items())[:4]:
                name = corpus.names.get(subj, subj)
                doc.add_paragraph(
                    f"For {per}, {name} recorded actual revenue of "
                    f"{format_money(actual, corpus.currency)} against a budget of "
                    f"{format_money(budget, corpus.currency)}.",
                    style="List Bullet",
                )
            continue

        if sec_idx == 5:
            table = doc.add_table(rows=1, cols=4)
            table.style = "Light Grid Accent 1"
            for cell, txt in zip(table.rows[0].cells, ["Risk", "Rating", "Owner", "Mitigation"]):
                cell.text = txt
            pairs = list(zip(corpus.risks, corpus.mitigations))
            sample_count = min(4, len(pairs))
            for (risk, rating), mit in rng.sample(pairs, sample_count):
                row = table.add_row().cells
                row[0].text = risk
                row[1].text = rating
                row[2].text = rng.choice(corpus.people)[0]
                row[3].text = mit

        doc.add_paragraph(f"{header_line} ({chart_title}) — {w.paragraph()}")
        for fig in sec_figs:
            fig_no += 1
            doc.add_paragraph(
                f"Visual Appendix {fig_no}: {header_line} — {chart_title} ({fig.caption}). "
                f"{w.sentence()}"
            )
            width = Inches(6.2) if fig.size[0] >= fig.size[1] else Inches(3.6)
            doc.add_picture(BytesIO(fig.data), width=width)
            cap = doc.add_paragraph()
            cap.add_run(f"Figure {fig_no}: {fig.caption}").bold = True
            doc.add_paragraph(w.paragraph(2))

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue(), text_unit_index


def _render_xlsx(
    rng: random.Random,
    *,
    corpus: ProseCorpus,
    header_line: str,
    governance_line: str,
    fact_sentence: str,
    chart_title: str,
    chart_value: str,
    target_bytes: int,
    theme: tuple[int, int, int] = (31, 119, 180),
    sheet_kind: str = "budget",
) -> tuple[bytes, int]:
    """Render a multi-sheet .xlsx workbook (README, Summary, Charts, and operational rows)."""
    figures = _build_encoded_figures(
        rng,
        corpus=corpus,
        header_line=header_line,
        chart_title=chart_title,
        chart_value=chart_value,
        target_bytes=target_bytes,
        theme=theme,
    )
    cols, make_row, _id_col, _probe_col, _probe_name, noun = sheet_definition(sheet_kind, corpus)
    buf = BytesIO()
    wb = xlsxwriter.Workbook(buf, {"in_memory": True})
    bold = wb.add_format({"bold": True})
    head = wb.add_format(
        {
            "bold": True,
            "bg_color": f"#{theme[0]:02x}{theme[1]:02x}{theme[2]:02x}",
            "font_color": "white",
        }
    )

    readme = wb.add_worksheet("README")
    readme.set_column(0, 0, 110)
    readme.write(0, 0, MARKING, bold)
    readme.write(1, 0, header_line, bold)
    readme.write(
        2,
        0,
        f"{corpus.company}. Extract date: 14 Oct 2026. {governance_line} Synthetic test data.",
    )
    readme.write(
        3,
        0,
        f"Sheets: Summary (revenue from agency charges), Charts ({chart_title}), "
        f"and operational data sheets for {noun}.",
    )
    readme.write(4, 0, f"Audit verification notes and financial overview for {header_line}:")
    readme.write(5, 0, f"Reference telemetry stream: {header_line} | {chart_title}")
    readme.write(6, 0, fact_sentence)
    for extra_idx in range(7, 15):
        readme.write(
            extra_idx,
            0,
            f"Section {extra_idx - 6} verification log for {header_line} — {chart_title}.",
        )
    text_unit_index = 7

    summary = wb.add_worksheet("Summary")
    summary.write_row(
        0,
        0,
        [
            "Line of work",
            "Period",
            f"Actual ({corpus.currency} k)",
            f"Budget ({corpus.currency} k)",
        ],
        head,
    )
    items = sorted(corpus.revenue.items())
    for r_idx, ((subject_id, period), (actual, budget)) in enumerate(items, 1):
        summary.write_row(
            r_idx,
            0,
            [corpus.names.get(subject_id, subject_id), period, actual, budget],
        )
    summary.set_column(0, 0, 44)
    if items:
        chart = wb.add_chart({"type": "column"})
        chart.add_series(
            {
                "name": "Actual",
                "categories": ["Summary", 1, 0, len(items), 0],
                "values": ["Summary", 1, 2, len(items), 2],
            }
        )
        chart.set_title({"name": "Revenue from agency charges"})
        summary.insert_chart("F2", chart, {"x_scale": 2.0, "y_scale": 1.5})

    charts_ws = wb.add_worksheet("Charts")
    row_at = 0
    for k, fig in enumerate(figures, 1):
        charts_ws.write(
            row_at,
            0,
            f"Visual Appendix {k}: {header_line} — {chart_title} ({fig.caption})",
            bold,
        )
        scale = 1100.0 / max(1, fig.size[0])
        charts_ws.insert_image(
            row_at + 1,
            0,
            f"figure{k}.png",
            {"image_data": BytesIO(fig.data), "x_scale": scale, "y_scale": scale},
        )
        row_at += int(fig.size[1] * scale / 20) + 4

    prefix = DATA_SHEET_PREFIX.get(sheet_kind, "Lines")
    data_ws = wb.add_worksheet(f"{prefix}_1")
    data_ws.write_row(0, 0, cols, head)
    data_ws.set_column(0, len(cols) - 1, 18)
    row_count = 100 if target_bytes >= 100_000 else 25
    for r_idx in range(1, row_count + 1):
        data_ws.write_row(r_idx, 0, make_row(rng, r_idx))

    wb.close()
    return buf.getvalue(), text_unit_index


def _render_pptx(
    rng: random.Random,
    *,
    corpus: ProseCorpus,
    header_line: str,
    governance_line: str,
    fact_sentence: str,
    chart_title: str,
    chart_value: str,
    target_bytes: int,
    theme: tuple[int, int, int] = (0, 90, 170),
) -> tuple[bytes, int]:
    """Render a 16:9 .pptx presentation with themed bands, speaker notes, and deep fact slide."""
    figures = _build_encoded_figures(
        rng,
        corpus=corpus,
        header_line=header_line,
        chart_title=chart_title,
        chart_value=chart_value,
        target_bytes=target_bytes,
        theme=theme,
    )
    prs = Presentation()
    prs.slide_width = PptxInches(13.333)
    prs.slide_height = PptxInches(7.5)
    theme_rgb = PptxRGBColor(*theme)
    w = Writer(rng, corpus)
    owner, role = rng.choice(corpus.people) if corpus.people else ("Tan Wei Jie", "Director")

    def _band(slide: object) -> None:
        bar = slide.shapes.add_shape(1, 0, 0, prs.slide_width, PptxInches(0.18))  # type: ignore[attr-defined]
        bar.fill.solid()
        bar.fill.fore_color.rgb = theme_rgb
        bar.line.fill.background()
        tb = slide.shapes.add_textbox(  # type: ignore[attr-defined]
            PptxInches(0.4), PptxInches(7.05), PptxInches(12.5), PptxInches(0.3)
        )
        tb.text_frame.text = f"{MARKING}   |   {header_line[:80]}"
        tb.text_frame.paragraphs[0].runs[0].font.size = PptxPt(10)

    def _bullets(frame: object, lines: Sequence[str], size: int = 18) -> None:
        frame.word_wrap = True  # type: ignore[attr-defined]
        frame.text = lines[0]  # type: ignore[attr-defined]
        for line in lines[1:]:
            frame.add_paragraph().text = line  # type: ignore[attr-defined]
        for p in frame.paragraphs:  # type: ignore[attr-defined]
            for r in p.runs:
                r.font.size = PptxPt(size)

    s_title = prs.slides.add_slide(prs.slide_layouts[0])
    s_title.shapes.title.text = header_line
    s_title.placeholders[1].text = (
        f"{owner}, {role}\n14 Oct 2026\n{governance_line}\nSynthetic test data"
    )
    _band(s_title)

    sections = _document_sections(corpus)
    s_agenda = prs.slides.add_slide(prs.slide_layouts[1])
    s_agenda.shapes.title.text = "Agenda"
    _bullets(s_agenda.placeholders[1].text_frame, sections, 20)
    _band(s_agenda)

    planned = _plan_section_figures(sections, figures)
    fig_no = 0
    text_unit_index = 6

    for heading, sec_figs in zip(sections, planned):
        if heading == "Financial overview":
            s_fin = prs.slides.add_slide(prs.slide_layouts[1])
            s_fin.shapes.title.text = f"{heading}: {header_line[:50]}"
            fin_lines = [fact_sentence, w.sentence(), w.sentence()]
            _bullets(s_fin.placeholders[1].text_frame, fin_lines, 16)
            _band(s_fin)
            text_unit_index = len(prs.slides)
            continue

        s_sec = prs.slides.add_slide(prs.slide_layouts[1])
        s_sec.shapes.title.text = heading
        _bullets(s_sec.placeholders[1].text_frame, [w.sentence() for _ in range(4)], 18)
        _band(s_sec)

        for fig in sec_figs:
            fig_no += 1
            s_fig = prs.slides.add_slide(prs.slide_layouts[5])
            s_fig.shapes.title.text = (
                f"Visual Appendix {fig_no}: {chart_title} ({fig.caption})"[:90]
            )
            s_fig.shapes.title.text_frame.paragraphs[0].runs[0].font.size = PptxPt(22)
            aspect = fig.size[0] / max(1, fig.size[1])
            box_w, box_h = (8.4, 5.2) if aspect >= 1 else (4.0, 5.4)
            pw, ph = (
                (box_w, box_w / aspect)
                if box_w / aspect <= box_h
                else (box_h * aspect, box_h)
            )
            s_fig.shapes.add_picture(
                BytesIO(fig.data),
                PptxInches(0.5),
                PptxInches(1.45),
                PptxInches(pw),
                PptxInches(ph),
            )
            tb = s_fig.shapes.add_textbox(
                PptxInches(0.9 + pw),
                PptxInches(1.6),
                PptxInches(max(2.0, 12.2 - pw)),
                PptxInches(5.0),
            )
            _bullets(tb.text_frame, [w.sentence() for _ in range(3)], 16)
            s_fig.notes_slide.notes_text_frame.text = (
                f"Figure {fig_no} ({header_line}). {w.paragraph(2)}"
            )
            _band(s_fig)

    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue(), text_unit_index


def _render_pdf(
    rng: random.Random,
    *,
    corpus: ProseCorpus,
    header_line: str,
    governance_line: str,
    fact_sentence: str,
    chart_title: str,
    chart_value: str,
    target_bytes: int,
    theme: tuple[int, int, int] = (0, 110, 110),
) -> tuple[bytes, int]:
    """Render an A4 Platypus PDF with page headers, financial table on page >= 6, and charts."""
    figures = _build_encoded_figures(
        rng,
        corpus=corpus,
        header_line=header_line,
        chart_title=chart_title,
        chart_value=chart_value,
        target_bytes=target_bytes,
        theme=theme,
    )
    styles = getSampleStyleSheet()
    pw, ph = A4
    frame_w, frame_h = pw - 4 * cm, ph - 5 * cm

    def _decorate(c: object, docobj: object) -> None:
        c.saveState()  # type: ignore[attr-defined]
        c.setFillColorRGB(theme[0] / 255.0, theme[1] / 255.0, theme[2] / 255.0)  # type: ignore[attr-defined]
        c.rect(0, ph - 0.5 * cm, pw, 0.5 * cm, fill=1, stroke=0)  # type: ignore[attr-defined]
        c.setFillColorRGB(0.65, 0, 0)  # type: ignore[attr-defined]
        c.setFont("Helvetica", 8)  # type: ignore[attr-defined]
        c.drawString(2 * cm, ph - 1.2 * cm, MARKING)  # type: ignore[attr-defined]
        c.setFillColorRGB(0.3, 0.3, 0.3)  # type: ignore[attr-defined]
        c.drawString(2 * cm, 1.2 * cm, header_line[:90])  # type: ignore[attr-defined]
        c.drawRightString(pw - 2 * cm, 1.2 * cm, f"Page {docobj.page}")  # type: ignore[attr-defined]
        c.restoreState()  # type: ignore[attr-defined]

    owner, role = rng.choice(corpus.people) if corpus.people else ("Tan Wei Jie", "Director")
    story: list[object] = [
        Spacer(1, 3 * cm),
        Paragraph(header_line, styles["Title"]),
        Paragraph(
            f"{corpus.company}<br/>Prepared by {owner}, {role}<br/>14 Oct 2026<br/>"
            f"{governance_line}",
            styles["Normal"],
        ),
        Spacer(1, 1 * cm),
        Paragraph(
            "<font color='#aa0000'>This document is synthetic test data. All "
            "names and figures are invented.</font>",
            styles["Normal"],
        ),
        PageBreak(),
    ]
    w = Writer(rng, corpus)
    sections = _document_sections(corpus)
    planned = _plan_section_figures(sections, figures)
    fig_no = 0
    text_unit_index = 6

    for heading, sec_figs in zip(sections, planned):
        story.append(Paragraph(heading.replace("&", "&amp;"), styles["Heading1"]))
        if heading == "Financial overview":
            story.append(
                Paragraph(
                    fact_sentence.replace("&", "&amp;").replace(": ", ":<br/>"),
                    styles["Normal"],
                )
            )
            rows = [
                [
                    "Line of work",
                    "Period",
                    f"Actual ({corpus.currency})",
                    f"Budget ({corpus.currency})",
                ]
            ]
            for (subj, per), (actual, budget) in sorted(corpus.revenue.items())[:6]:
                rows.append(
                    [
                        corpus.names.get(subj, subj)[:40],
                        per,
                        format_money(actual, corpus.currency),
                        format_money(budget, corpus.currency),
                    ]
                )
            tbl = Table(rows, hAlign="LEFT")
            tbl.setStyle(
                TableStyle(
                    [
                        ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                        ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                        ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ]
                )
            )
            story += [Spacer(1, 0.3 * cm), tbl, PageBreak()]
            continue

        story.append(Paragraph(w.paragraph(), styles["Normal"]))
        for fig in sec_figs:
            fig_no += 1
            aspect = fig.size[0] / max(1, fig.size[1])
            iw = frame_w if aspect >= 1 else frame_w * 0.55
            ih = iw / aspect
            if ih > frame_h * 0.65:
                ih = frame_h * 0.65
                iw = ih * aspect
            story += [
                Spacer(1, 0.25 * cm),
                RLImage(BytesIO(fig.data), width=iw, height=ih),
                Paragraph(
                    f"<b>Visual Appendix {fig_no}: {header_line} — {chart_title} "
                    f"({fig.caption})</b>",
                    styles["Italic"],
                ),
                Paragraph(w.paragraph(2), styles["Normal"]),
            ]
        story.append(PageBreak())

    buf = BytesIO()
    SimpleDocTemplate(
        buf,
        pagesize=A4,
        title=header_line,
        author=owner,
        leftMargin=2 * cm,
        rightMargin=2 * cm,
        topMargin=2.2 * cm,
        bottomMargin=2.2 * cm,
    ).build(story, onFirstPage=_decorate, onLaterPages=_decorate)
    return buf.getvalue(), text_unit_index


def _render_by_format(
    fmt: Literal["docx", "xlsx", "pptx", "pdf"],
    rng: random.Random,
    *,
    corpus: ProseCorpus,
    header_line: str,
    governance_line: str,
    fact_sentence: str,
    chart_title: str,
    chart_value: str,
    target_bytes: int,
    theme: tuple[int, int, int] = (0, 84, 147),
    sheet_kind: str = "budget",
) -> tuple[bytes, int]:
    if fmt == "docx":
        return _render_docx(
            rng,
            corpus=corpus,
            header_line=header_line,
            governance_line=governance_line,
            fact_sentence=fact_sentence,
            chart_title=chart_title,
            chart_value=chart_value,
            target_bytes=target_bytes,
            theme=theme,
        )
    if fmt == "xlsx":
        return _render_xlsx(
            rng,
            corpus=corpus,
            header_line=header_line,
            governance_line=governance_line,
            fact_sentence=fact_sentence,
            chart_title=chart_title,
            chart_value=chart_value,
            target_bytes=target_bytes,
            theme=theme,
            sheet_kind=sheet_kind,
        )
    if fmt == "pptx":
        return _render_pptx(
            rng,
            corpus=corpus,
            header_line=header_line,
            governance_line=governance_line,
            fact_sentence=fact_sentence,
            chart_title=chart_title,
            chart_value=chart_value,
            target_bytes=target_bytes,
            theme=theme,
        )
    if fmt == "pdf":
        return _render_pdf(
            rng,
            corpus=corpus,
            header_line=header_line,
            governance_line=governance_line,
            fact_sentence=fact_sentence,
            chart_title=chart_title,
            chart_value=chart_value,
            target_bytes=target_bytes,
            theme=theme,
        )
    raise ValueError(f"Unsupported format: {fmt}")


def _unique_currency(rng: random.Random, used: set[str], currency: str = "SGD") -> str:
    while True:
        val = f"{currency} {rng.randint(1, 9)},{rng.randint(100, 999)},{rng.randint(100, 999)}"
        if val not in used:
            used.add(val)
            return val


def _unique_percentage(rng: random.Random, used: set[str]) -> str:
    while True:
        val = f"{rng.randint(11, 89)}.{rng.randint(11, 99)}%"
        if val not in used:
            used.add(val)
            return val


def build_corpus(
    out_dir: Path,
    *,
    sizes_mb: Sequence[float] = (18.0, 50.0, 100.0),
    seed: int = 42,
    formats: Sequence[str] = ("docx", "xlsx", "pptx", "pdf"),
    corpus_dir: Path | None = None,
    world_pack: str | None = None,
    domain: str | None = None,
    locale: str | None = None,
    facets: Mapping[str, str] | None = None,
    subject: str | None = None,
    writer: str = "template",
) -> list[ManifestEntry]:
    """Build large native documents (up to 1 GiB) and paired < 1 MB distractor trap files."""
    if not sizes_mb or any(s <= 0 or s > 1024.0 for s in sizes_mb):
        raise ValueError("Every entry in sizes_mb must be between 0 and 1024 MiB (1 GiB).")

    rng = random.Random(seed)
    ctx = _resolve_world_context(
        seed=seed,
        out_dir=out_dir,
        corpus_dir=corpus_dir,
        world_pack=world_pack,
        domain=domain,
        locale=locale,
        facets=facets,
        writer=writer,
    )
    topics = (
        tuple(f"{t} — {subject.strip()}" for t in ctx.topics)
        if subject and subject.strip()
        else ctx.topics
    )
    out_dir.mkdir(parents=True, exist_ok=True)
    (out_dir / "large").mkdir(parents=True, exist_ok=True)
    (out_dir / "traps").mkdir(parents=True, exist_ok=True)

    used_values: set[str] = set()
    entries: list[ManifestEntry] = []
    item_counter = 0

    for size_mb in sizes_mb:
        min_target_bytes = int(size_mb * 1024 * 1024)
        size_tag = f"{size_mb:g}mb".replace(".", "_")
        for fmt_raw in formats:
            fmt = fmt_raw.lower().strip().removeprefix(".")
            if fmt not in ("docx", "xlsx", "pptx", "pdf"):
                raise ValueError(f"Unsupported format '{fmt_raw}'.")
            typed_fmt: Literal["docx", "xlsx", "pptx", "pdf"] = fmt  # type: ignore[assignment]

            base_topic = topics[item_counter % len(topics)]
            codename = ctx.prose_corpus.brand_tokens[
                item_counter % len(ctx.prose_corpus.brand_tokens)
            ]
            theme = THEME_PALETTE[item_counter % len(THEME_PALETTE)]
            sheet_kind = SHEET_KINDS[item_counter % len(SHEET_KINDS)]
            item_counter += 1
            file_id = f"bl-{item_counter:03d}-{fmt}-{size_tag}"
            if codename not in base_topic:
                topic = f"{base_topic} ({codename}) Unit-{item_counter:03d}"
            else:
                topic = f"{base_topic} Unit-{item_counter:03d}"
            period = "FY2026"
            distractor_period = "FY2025"

            text_metric_label = f"{topic} {period} Authorized Capital Reserve"
            image_chart_title = f"{topic} {period} Q4 Peak Thermal Efficiency"

            text_golden = _unique_currency(rng, used_values, currency=ctx.currency)
            text_canary = _unique_currency(rng, used_values, currency=ctx.currency)
            image_golden = _unique_percentage(rng, used_values)
            image_canary = _unique_percentage(rng, used_values)

            target_rel = f"large/{file_id}_fy2026_final.{fmt}"
            distractor_rel = f"traps/{file_id}_fy2025_draft.{fmt}"

            target_header = (
                f"{topic} {period} Final Audit Report | "
                f"{text_metric_label} | {image_chart_title}"
            )
            target_fact = (
                f"Verified {text_metric_label} for {topic} ({period} Final): {text_golden}."
            )
            target_bytes, text_unit_idx = _render_by_format(
                typed_fmt,
                rng,
                corpus=ctx.prose_corpus,
                header_line=target_header,
                governance_line=ctx.governance_line,
                fact_sentence=target_fact,
                chart_title=image_chart_title,
                chart_value=image_golden,
                target_bytes=min_target_bytes,
                theme=theme,
                sheet_kind=sheet_kind,
            )

            dist_header = f"{topic} {distractor_period} Preliminary Draft Report"
            dist_fact = (
                f"Preliminary {topic} {distractor_period} Draft Capital Reserve: "
                f"{text_canary} (Chart Draft Metric: {image_canary})."
            )
            dist_bytes, _ = _render_by_format(
                typed_fmt,
                rng,
                corpus=ctx.prose_corpus,
                header_line=dist_header,
                governance_line=ctx.governance_line,
                fact_sentence=dist_fact,
                chart_title=f"{topic} {distractor_period} Draft Chart",
                chart_value=image_canary,
                target_bytes=24_576,
                theme=theme,
                sheet_kind=sheet_kind,
            )

            target_path = out_dir / target_rel
            dist_path = out_dir / distractor_rel
            target_path.write_bytes(target_bytes)
            dist_path.write_bytes(dist_bytes)

            entries.append(
                ManifestEntry(
                    file_id=file_id,
                    format=typed_fmt,
                    target_file=target_rel,
                    target_size_bytes=len(target_bytes),
                    target_sha256=hashlib.sha256(target_bytes).hexdigest(),
                    distractor_file=distractor_rel,
                    distractor_size_bytes=len(dist_bytes),
                    distractor_sha256=hashlib.sha256(dist_bytes).hexdigest(),
                    topic=topic,
                    period=period,
                    distractor_period=distractor_period,
                    text_metric_label=text_metric_label,
                    text_golden_value=text_golden,
                    text_canary_value=text_canary,
                    text_unit_index=text_unit_idx,
                    image_chart_title=image_chart_title,
                    image_golden_value=image_golden,
                    image_canary_value=image_canary,
                    worldloom_commit=ctx.provenance.worldloom_commit,
                    world_seed=ctx.provenance.world_seed,
                    recipe_digest=ctx.provenance.recipe_digest,
                    narration=ctx.provenance.narration,
                    model_id=ctx.provenance.model_id,
                    prompt_version=ctx.provenance.prompt_version,
                )
            )

    manifest_path = out_dir / "manifest.jsonl"
    manifest_lines = [entry.model_dump_json() for entry in entries]
    manifest_path.write_text("\n".join(manifest_lines) + "\n", encoding="utf-8")
    generate_queries(manifest_path, out_dir / "cases.json")
    return entries
