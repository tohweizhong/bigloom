"""Add BigLoom test evidence to WorldLoom documents and files.

BigLoom writes no prose. It adds three things to a WorldLoom document:

1. A text fact: a small table section in a copy of the IR, before rendering.
2. An image fact: a chart picture in the rendered file, after rendering. The value is in
   the picture only, never in the file text.
3. Byte padding: an unreferenced, stored ZIP member for OOXML files, or an embedded file
   for PDF files, so that each file reaches its target size.

It also fixes broken placeholder sizes in the WorldLoom enterprise deck template.
"""

from __future__ import annotations

import random
import zipfile
from collections.abc import Sequence
from io import BytesIO
from typing import TYPE_CHECKING

from PIL import Image, ImageDraw, ImageFont

if TYPE_CHECKING:  # pragma: no cover
    from worldloom.models import ArtifactIR, ArtifactSection

EVIDENCE_KEY = "bigloom-evidence"
BALLAST_NAME = "customXml/bigloom-ballast.bin"
PDF_BALLAST_NAME = "bigloom-ballast.bin"
_CHUNK = 1 << 20


# ---------------------------------------------------------------------------
# Text fact in the IR
# ---------------------------------------------------------------------------


def evidence_section(heading: str, rows: Sequence[tuple[str, str]]) -> ArtifactSection:
    """Return one table section with (measure, value) rows."""
    from worldloom.models import ArtifactSection, Cell, Column, Row, Table

    table = Table(
        key=EVIDENCE_KEY,
        title=heading,
        columns=[Column(key="measure", label="Measure"), Column(key="value", label="Value")],
        rows=[
            Row(key=f"r{i}", label=label, cells={"measure": Cell(value=label), "value": Cell(value=value)})
            for i, (label, value) in enumerate(rows, 1)
        ],
    )
    return ArtifactSection(heading=heading, table=table)


def with_evidence(ir: ArtifactIR, section: ArtifactSection, *, depth: int = 6, title_suffix: str = "") -> ArtifactIR:
    """Return a copy of the IR with the section inserted at position `depth`, or at the end."""
    sections = list(ir.sections)
    sections.insert(min(depth, len(sections)), section)
    update: dict[str, object] = {"sections": sections}
    if title_suffix:
        update["title"] = f"{ir.title}{title_suffix}"
    return ir.model_copy(update=update)


# ---------------------------------------------------------------------------
# Image fact
# ---------------------------------------------------------------------------


def _font(size: int) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    for name in ("DejaVuSans.ttf", "Arial.ttf", "LiberationSans-Regular.ttf"):
        try:
            return ImageFont.truetype(name, size)
        except OSError:
            continue
    return ImageFont.load_default()


def fact_chart_png(title: str, categories: Sequence[str], value_label: str, seed: int, peak: int = 0) -> bytes:
    """Draw a bar chart. The bar at index `peak` is the tallest and carries `value_label`.

    `value_label` is the image fact. Other bar heights come from a seeded generator, so the
    same seed draws the same chart.
    """
    if not categories:
        raise ValueError("A chart needs at least one category.")
    rng = random.Random(seed)
    cats = [c[:18] for c in categories][:6]
    peak = peak % len(cats)
    heights = [rng.randint(180, 520) for _ in cats]
    heights[peak] = 640
    img = Image.new("RGB", (1600, 1000), "white")
    draw = ImageDraw.Draw(img)
    draw.text((80, 40), title[:80], fill=(20, 20, 20), font=_font(40))
    base_y = 860
    slot = 1400 // len(cats)
    for i, (cat, h) in enumerate(zip(cats, heights, strict=True)):
        x0 = 120 + i * slot + slot // 6
        x1 = x0 + slot * 2 // 3
        draw.rectangle([x0, base_y - h, x1, base_y], fill=(0, 84, 147) if i != peak else (200, 16, 46))
        draw.text((x0, base_y + 16), cat, fill=(40, 40, 40), font=_font(26))
        if i == peak:
            draw.text((x0, base_y - h - 56), value_label, fill=(20, 20, 20), font=_font(40))
    draw.line([100, base_y, 1520, base_y], fill=(0, 0, 0), width=4)
    buf = BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


def docx_add_picture(data: bytes, png: bytes, caption: str, *, at: float = 0.6) -> bytes:
    """Insert a picture and its caption at fraction `at` of the body paragraphs."""
    from docx import Document
    from docx.shared import Inches

    doc = Document(BytesIO(data))
    paras = doc.paragraphs
    anchor = paras[min(len(paras) - 1, max(0, int(len(paras) * at)))]
    pic = doc.add_paragraph()
    pic.add_run().add_picture(BytesIO(png), width=Inches(6.0))
    cap = doc.add_paragraph(caption)
    anchor._p.addprevious(pic._p)
    anchor._p.addprevious(cap._p)
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def pdf_add_picture(data: bytes, png: bytes, caption: str) -> bytes:
    """Insert one A4 page with the picture and its caption at the middle of the document."""
    from pypdf import PdfReader, PdfWriter
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    page_buf = BytesIO()
    c = canvas.Canvas(page_buf, pagesize=A4, invariant=True)
    w, h = A4
    c.drawImage(ImageReader(BytesIO(png)), 50, h / 2 - 150, width=w - 100, height=(w - 100) * 10 / 16)
    c.setFont("Helvetica", 10)
    c.drawString(50, h / 2 - 170, caption[:110])
    c.save()
    reader = PdfReader(BytesIO(data))
    extra = PdfReader(BytesIO(page_buf.getvalue())).pages[0]
    writer = PdfWriter()
    mid = len(reader.pages) // 2
    for i, page in enumerate(reader.pages):
        if i == mid:
            writer.add_page(extra)
        writer.add_page(page)
    if reader.metadata:
        writer.add_metadata({k: str(v) for k, v in reader.metadata.items()})
    out = BytesIO()
    writer.write(out)
    return out.getvalue()


def pptx_add_evidence_slide(data: bytes, title: str, lines: Sequence[str], png: bytes | None) -> bytes:
    """Add one slide with text lines and an optional picture, at the middle of the deck."""
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation(BytesIO(data))
    layout = prs.slide_layouts[5] if len(prs.slide_layouts) > 5 else prs.slide_layouts[-1]
    slide = prs.slides.add_slide(layout)
    if slide.shapes.title is not None:
        slide.shapes.title.text = title
    box = slide.shapes.add_textbox(Inches(0.6), Inches(1.6), Inches(5.6), Inches(3.0))
    frame = box.text_frame
    frame.word_wrap = True
    for i, line in enumerate(lines):
        para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        para.text = line
        for run in para.runs:
            run.font.size = Pt(18)
    if png is not None:
        slide.shapes.add_picture(BytesIO(png), Inches(6.5), Inches(1.6), width=Inches(6.2))
    ids = prs.slides._sldIdLst
    new = ids[-1]
    ids.remove(new)
    ids.insert(len(ids) // 2, new)
    out = BytesIO()
    prs.save(out)
    return out.getvalue()


_SHEET_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
    '<sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>{caption}</t></is></c></row></sheetData>'
    '<drawing r:id="rId1"/></worksheet>'
)
_DRAWING_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    '<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
    '<xdr:oneCellAnchor><xdr:from><xdr:col>0</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>2</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from>'
    '<xdr:ext cx="7315200" cy="4572000"/><xdr:pic><xdr:nvPicPr><xdr:cNvPr id="2" name="Chart"/><xdr:cNvPicPr/></xdr:nvPicPr>'
    '<xdr:blipFill><a:blip r:embed="rId1"/><a:stretch><a:fillRect/></a:stretch></xdr:blipFill>'
    '<xdr:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="7315200" cy="4572000"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></xdr:spPr>'
    '</xdr:pic><xdr:clientData/></xdr:oneCellAnchor></xdr:wsDr>'
)
_REL = '<Relationship Id="{rid}" Type="{type}" Target="{target}"/>'
_RELS = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">{body}</Relationships>'
)
_T_SHEET = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet"
_T_DRAWING = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing"
_T_IMAGE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"


def xlsx_add_picture_sheet(data: bytes, png: bytes, sheet_name: str, caption: str) -> bytes:
    """Add a worksheet that shows a picture, by editing the package parts.

    openpyxl drops native charts when it saves a workbook it loaded, so this edits the ZIP
    parts directly and leaves every WorldLoom part as it was.
    """
    import re
    from xml.sax.saxutils import escape

    src = zipfile.ZipFile(BytesIO(data))
    names = src.namelist()
    sheet_nums = [int(m.group(1)) for n in names if (m := re.match(r"xl/worksheets/sheet(\d+)\.xml$", n))]
    sheet_n = max(sheet_nums, default=0) + 1
    drawing_nums = [int(m.group(1)) for n in names if (m := re.match(r"xl/drawings/drawing(\d+)\.xml$", n))]
    drawing_n = max(drawing_nums, default=0) + 1
    image_n = sum(1 for n in names if n.startswith("xl/media/")) + 1
    workbook = src.read("xl/workbook.xml").decode("utf-8")
    wb_rels = src.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    types = src.read("[Content_Types].xml").decode("utf-8")
    rid_nums = [int(x) for x in re.findall(r'Id="rId(\d+)"', wb_rels)]
    rid = f"rId{max(rid_nums, default=0) + 1}"
    sheet_ids = [int(x) for x in re.findall(r'sheetId="(\d+)"', workbook)]
    safe_name = escape(sheet_name[:31]).replace('"', "&quot;")
    root_tag = re.search(r"<workbook\b[^>]*>", workbook)
    if root_tag and "xmlns:r=" not in root_tag.group(0):
        # Declare the relationships prefix on the root element when the writer left it out.
        workbook = re.sub(
            r"<workbook\b",
            '<workbook xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"',
            workbook,
            count=1,
        )
    workbook = workbook.replace(
        "</sheets>", f'<sheet name="{safe_name}" sheetId="{max(sheet_ids, default=0) + 1}" r:id="{rid}"/></sheets>'
    )
    wb_rels = wb_rels.replace(
        "</Relationships>", _REL.format(rid=rid, type=_T_SHEET, target=f"worksheets/sheet{sheet_n}.xml") + "</Relationships>"
    )
    overrides = (
        f'<Override PartName="/xl/worksheets/sheet{sheet_n}.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
        f'<Override PartName="/xl/drawings/drawing{drawing_n}.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.drawing+xml"/>'
    )
    if 'Extension="png"' not in types:
        overrides += '<Default Extension="png" ContentType="image/png"/>'
    types = types.replace("</Types>", overrides + "</Types>")
    out = BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            if info.filename == "xl/workbook.xml":
                dst.writestr(info, workbook)
            elif info.filename == "xl/_rels/workbook.xml.rels":
                dst.writestr(info, wb_rels)
            elif info.filename == "[Content_Types].xml":
                dst.writestr(info, types)
            else:
                dst.writestr(info, src.read(info.filename))
        dst.writestr(f"xl/worksheets/sheet{sheet_n}.xml", _SHEET_XML.format(caption=escape(caption)))
        dst.writestr(
            f"xl/worksheets/_rels/sheet{sheet_n}.xml.rels",
            _RELS.format(body=_REL.format(rid="rId1", type=_T_DRAWING, target=f"../drawings/drawing{drawing_n}.xml")),
        )
        dst.writestr(f"xl/drawings/drawing{drawing_n}.xml", _DRAWING_XML)
        dst.writestr(
            f"xl/drawings/_rels/drawing{drawing_n}.xml.rels",
            _RELS.format(body=_REL.format(rid="rId1", type=_T_IMAGE, target=f"../media/image{image_n}.png")),
        )
        dst.writestr(f"xl/media/image{image_n}.png", png)
    return out.getvalue()


# ---------------------------------------------------------------------------
# Deck layout fix
# ---------------------------------------------------------------------------


def fix_deck_layout(data: bytes) -> tuple[bytes, int]:
    """Give a valid size to every layout and slide placeholder with no height or wider than the slide.

    The WorldLoom enterprise deck template gives title and body placeholders a width of
    16 in and a height of 0 on a 13.33 in slide, so PowerPoint draws titles off the slide.
    Return the new bytes and the number of placeholders changed.
    """
    from pptx import Presentation
    from pptx.enum.shapes import PP_PLACEHOLDER
    from pptx.util import Inches

    prs = Presentation(BytesIO(data))
    w, h = prs.slide_width, prs.slide_height
    margin = Inches(0.6)
    inner = w - 2 * margin
    titles = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}
    changed = 0
    containers = [*prs.slide_layouts, *prs.slides]
    for container in containers:
        for ph in container.placeholders:
            if ph.height and ph.width and ph.left is not None and ph.left + ph.width <= w:
                continue
            kind = ph.placeholder_format.type
            if kind in titles:
                geometry = (margin, Inches(0.35), inner, Inches(1.0))
            elif kind == PP_PLACEHOLDER.DATE:
                geometry = (margin, h - Inches(0.55), Inches(3.0), Inches(0.35))
            elif kind == PP_PLACEHOLDER.FOOTER:
                geometry = ((w - Inches(6.0)) // 2, h - Inches(0.55), Inches(6.0), Inches(0.35))
            elif kind == PP_PLACEHOLDER.SLIDE_NUMBER:
                geometry = (w - margin - Inches(1.5), h - Inches(0.55), Inches(1.5), Inches(0.35))
            else:
                geometry = (margin, Inches(1.5), inner, h - Inches(2.2))
            ph.left, ph.top, ph.width, ph.height = geometry
            changed += 1
    out = BytesIO()
    prs.save(out)
    return out.getvalue(), changed


# ---------------------------------------------------------------------------
# Byte padding
# ---------------------------------------------------------------------------


def _random_chunks(total: int, seed: int):  # type: ignore[no-untyped-def]
    rng = random.Random(seed)
    left = total
    while left > 0:
        step = min(left, _CHUNK)
        yield rng.randbytes(step)
        left -= step


def _pad_ooxml_once(data: bytes, need: int, seed: int) -> bytes:
    buf = BytesIO(data)
    with zipfile.ZipFile(buf, "a", zipfile.ZIP_STORED, allowZip64=True) as zf:
        info = zipfile.ZipInfo(BALLAST_NAME, date_time=(2026, 1, 1, 0, 0, 0))
        with zf.open(info, "w", force_zip64=need > (1 << 31)) as fh:
            for chunk in _random_chunks(need, seed):
                fh.write(chunk)
    return buf.getvalue()


def pad_ooxml(data: bytes, target_bytes: int, seed: int) -> bytes:
    """Add an unreferenced, stored ZIP member so that the package reaches `target_bytes`.

    Office applications ignore a part that no relationship names. The member holds seeded
    random bytes, so it does not compress and the same seed gives the same file. The result
    is at least `target_bytes` long and at most a few hundred bytes longer.
    """
    if len(data) >= target_bytes:
        return data
    out = _pad_ooxml_once(data, target_bytes - len(data), seed)
    overhead = len(out) - target_bytes
    if overhead > 0:
        out = _pad_ooxml_once(data, max(1, target_bytes - len(data) - overhead), seed)
    return out


def _pad_pdf_once(data: bytes, need: int, seed: int) -> bytes:
    from pypdf import PdfReader, PdfWriter

    writer = PdfWriter(clone_from=PdfReader(BytesIO(data)))
    writer.add_attachment(PDF_BALLAST_NAME, b"".join(_random_chunks(need, seed)))
    out = BytesIO()
    writer.write(out)
    return out.getvalue()


def pad_pdf(data: bytes, target_bytes: int, seed: int) -> bytes:
    """Embed a file of seeded random bytes so that the PDF reaches at least `target_bytes`."""
    if len(data) >= target_bytes:
        return data
    need = target_bytes - len(data)
    out = _pad_pdf_once(data, need, seed)
    gap = target_bytes - len(out)
    if gap != 0:
        # The rewrite changes the size of the original objects. Correct the payload once.
        out = _pad_pdf_once(data, max(1, need + gap), seed)
    if len(out) < target_bytes:
        out = _pad_pdf_once(data, need + (target_bytes - len(out)) + 64, seed)
    return out
