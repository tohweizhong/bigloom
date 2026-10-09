"""Assemble large .docx proposals, < 1 MB trap twins, manifests, and evaluation files."""

from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import random
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Literal

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Inches, Pt, RGBColor
from pydantic import BaseModel, ConfigDict
from worldloom.native_artifacts import inspect_artifact

from .facts import (
    ProposalTextFact,
    WorldEntities,
    load_world_entities,
    select_proposal_facts,
)
from .plates import Fig, encode_all, make_figs, make_trap_figs
from .prose import ProseClientProtocol, load_prose, write_prose
from .spec import ProposalSpec
from .verify import find_leaks


class DocxLocator(BaseModel):
    """Exact location of a text or image fact inside a rendered .docx archive."""

    model_config = ConfigDict(frozen=True, extra="forbid")

    kind: Literal["paragraph", "image"]
    unit_locator: str
    index: int
    caption: str | None = None


@dataclass(frozen=True)
class ProposalBuildResult:
    """Paths and metadata produced by a proposal build."""

    manifest_entries: list[dict[str, object]]
    eval_rows: list[dict[str, object]]
    target_paths: list[Path]
    trap_path: Path


def _size_label(size_mb: int) -> str:
    if size_mb >= 1024 and size_mb % 1024 == 0:
        return f"{size_mb // 1024}gb"
    return f"{size_mb}mb"


def _add_prose(doc: object, text: str) -> int:
    words = 0
    for block in re.split(r"\n\s*\n", text):
        cleaned = re.sub(r"[*#`]", "", block).strip()
        if not cleaned:
            continue
        lines = cleaned.splitlines()
        if all(line.strip().startswith(("- ", "• ")) for line in lines):
            for line in lines:
                doc.add_paragraph(line.strip()[2:].strip(), style="List Bullet")  # type: ignore[attr-defined]
                words += len(line.split())
        else:
            doc.add_paragraph(" ".join(line.strip() for line in lines))  # type: ignore[attr-defined]
            words += len(cleaned.split())
    return words


def _add_table(doc: object, header_row: list[str], rows: list[list[str]]) -> None:
    t = doc.add_table(rows=1, cols=len(header_row))  # type: ignore[attr-defined]
    t.style = "Light Grid Accent 1"
    for cell, txt in zip(t.rows[0].cells, header_row):
        cell.text = txt
    for row in rows:
        for cell, txt in zip(t.add_row().cells, row):
            cell.text = str(txt)


def write_proposal_docx(
    path: Path,
    spec: ProposalSpec,
    entities: WorldEntities,
    prose: list[tuple[str, list[tuple[str, str]]]],
    facts: list[ProposalTextFact],
    figs: list[Fig],
    rng: random.Random,
    *,
    header_filename: str | None = None,
    trap: bool = False,
) -> dict[str, int]:
    """Render a target or trap .docx proposal to `path`."""
    doc = Document()
    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Calibri"
    normal_style.font.size = Pt(11)

    display_name = header_filename or path.name
    doc.sections[0].header.paragraphs[0].text = entities.marking
    doc.sections[0].footer.paragraphs[0].text = f"{display_name}  |  {entities.agency_name}"
    doc.add_heading(spec.title, 0)

    status = (
        "DRAFT v0.3, 14 Feb 2026. SUPERSEDED by v1.0. Do not use."
        if trap
        else "Version 1.0 (FINAL), 28 Sep 2026."
    )
    doc.add_paragraph().add_run(f"{entities.agency_name}. {status}").italic = True
    red = doc.add_paragraph().add_run(
        "This document is synthetic test data. All names, figures, and decisions are invented."
    )
    red.font.color.rgb = RGBColor(0xAA, 0, 0)
    _add_table(
        doc,
        ["Version", "Date", "Change"],
        [
            ["0.1", "9 Jan 2026", "First outline"],
            ["0.3", "14 Feb 2026", "Agency review"],
            ["1.0", "28 Sep 2026", "Final for approval"],
        ][: 2 if trap else 3],
    )

    def sentence(f: ProposalTextFact) -> str:
        return f.trap_sentence if trap else f.sentence

    def pick_person() -> str:
        return f"{rng.choice(entities.given_names)} {rng.choice(entities.surnames)}"

    words = 0
    fig_i = 0
    total_sections = max(1, sum(len(s) for _, s in prose))
    per_section = max(1, math.ceil(len(figs) / total_sections))

    for chapter, secs in prose:
        doc.add_heading(chapter, level=1)
        if chapter == "Executive summary":
            doc.add_paragraph("Key decisions sought:", style="List Bullet")
            for f in facts:
                if f.where == "start":
                    doc.add_paragraph(sentence(f), style="List Bullet")
        for si, (title, text) in enumerate(secs):
            doc.add_heading(title, level=2)
            words += _add_prose(doc, text)
            if chapter == spec.middle_chapter and si == len(secs) // 2:
                doc.add_paragraph("Phase 1 delivery commitments:")
                for f in facts:
                    if f.where == "middle":
                        doc.add_paragraph(sentence(f), style="List Bullet")
            if chapter == "Financial analysis and budget" and si == 0:
                _add_table(
                    doc,
                    ["Work package", "FY2026", "FY2027", "FY2028", "FY2029"],
                    [
                        [p]
                        + [
                            f"{entities.currency} {rng.randint(2, 30)}.{rng.randint(0, 9)}m"
                            for _ in range(4)
                        ]
                        for p in entities.packages
                    ],
                )
            if chapter == "Risk management" and si == 0:
                risks = [
                    "Sensor supply delays",
                    "Data sharing agreements not signed in time",
                    "Site works change precinct access",
                    "Cloud cost overrun",
                    "Vendor capacity",
                    "Cyber security findings before go-live",
                ]
                _add_table(
                    doc,
                    ["Risk", "Likelihood", "Impact", "Owner"],
                    [
                        [
                            r,
                            rng.choice(["Low", "Medium", "High"]),
                            rng.choice(["Medium", "High"]),
                            pick_person(),
                        ]
                        for r in risks
                    ],
                )
            for _ in range(per_section):
                if fig_i >= len(figs):
                    break
                fig = figs[fig_i]
                fig_i += 1
                doc.add_paragraph(
                    f"Plate {fig_i} shows the {fig.caption[0].lower()}{fig.caption[1:]}."
                )
                doc.add_picture(io.BytesIO(fig.data), width=Inches(6.3))
                doc.add_paragraph().add_run(f"Plate {fig_i}: {fig.caption}").bold = True

    while fig_i < len(figs):
        fig = figs[fig_i]
        fig_i += 1
        doc.add_paragraph(f"Plate {fig_i} shows the {fig.caption[0].lower()}{fig.caption[1:]}.")
        doc.add_picture(io.BytesIO(fig.data), width=Inches(6.3))
        doc.add_paragraph().add_run(f"Plate {fig_i}: {fig.caption}").bold = True

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading(spec.approval_annex, level=1)
    doc.add_paragraph(
        "The approving authority considered this proposal at its meeting and "
        "recorded the decision below."
    )
    for f in facts:
        if f.where == "end":
            doc.add_paragraph(sentence(f))
    doc.add_paragraph("End of document.")
    doc.save(path)
    return {"words": words, "figures": fig_i}


def _resolve_text_locator(payload: bytes, needle: str) -> DocxLocator:
    snapshot = inspect_artifact(payload, "docx")
    for idx, unit in enumerate(snapshot.units, start=1):
        if needle in unit.text:
            return DocxLocator(
                kind="paragraph",
                unit_locator=unit.locator,
                index=idx,
            )
    raise ValueError(f"Unable to locate text fact {needle!r} in rendered .docx.")


def build_proposal(
    spec: ProposalSpec,
    out_dir: Path,
    *,
    seed: int = 2026,
    skip_prose: bool = False,
    prose_client: ProseClientProtocol | None = None,
) -> ProposalBuildResult:
    """Build all target .docx sizes, the < 1 MB trap twin, manifest.jsonl, and evals files."""
    (out_dir / "docx").mkdir(parents=True, exist_ok=True)
    (out_dir / "traps").mkdir(parents=True, exist_ok=True)
    cache_dir = out_dir / "_prose_cache"

    entities = load_world_entities(spec, seed=seed)
    if not skip_prose:
        write_prose(cache_dir, spec, entities, client=prose_client, seed=seed)
    prose = load_prose(cache_dir, spec)

    rng = random.Random(seed)
    facts = select_proposal_facts(spec, entities, seed=seed)
    figs = make_figs(rng, spec.figures, spec, entities)
    picks = {"start": 0, "middle": len(figs) // 2, "end": len(figs) - 1}

    if spec.leak_terms:
        texts_to_check = [
            f"{spec.filename} {spec.trap_filename}",
            (
                f"{spec.title} {spec.short_title} {spec.doc_ref} {spec.office} "
                f"{entities.agency_name} {entities.marking} {entities.short_mark} "
                f"{spec.approval_annex}"
            ),
            " ".join(
                list(spec.chapters)
                + list(entities.zones)
                + list(entities.packages)
                + list(entities.phases)
            ),
            json.dumps([f.__dict__ for f in facts]),
            json.dumps([[f.caption, f.question, f.answer] for f in figs]),
        ]
        for chapter, secs in prose:
            for title, body in secs:
                texts_to_check.append(f"{chapter} {title} {body}")
        leaked = find_leaks(texts_to_check, spec.leak_terms)
        if leaked:
            raise ValueError(f"Forbidden leak_terms found in proposal inputs: {leaked}")

    # Build the < 1 MB trap twin first so its SHA-256 and byte size can be recorded in manifest.jsonl
    trap_figs = make_trap_figs(figs, picks, spec, entities, seed=seed, budget=80_000)
    second_chapter = spec.chapters[1] if len(spec.chapters) > 1 else spec.chapters[0]
    trap_prose = [
        (c, s[:1])
        for c, s in prose
        if c in ("Executive summary", second_chapter, spec.middle_chapter)
    ]
    trap_path = out_dir / "traps" / spec.trap_filename
    write_proposal_docx(
        trap_path,
        spec,
        entities,
        trap_prose,
        facts,
        [f for _, f in trap_figs],
        random.Random(7),
        header_filename=spec.trap_filename,
        trap=True,
    )
    trap_bytes = trap_path.read_bytes()
    trap_sha256 = hashlib.sha256(trap_bytes).hexdigest()

    # Build each requested size in sizes_mb
    target_paths: list[Path] = []
    manifest_entries: list[dict[str, object]] = []
    stem = Path(spec.filename).stem
    suffix = Path(spec.filename).suffix or ".docx"

    for size_mb in spec.sizes_mb:
        target_bytes = int(size_mb) * 1024 * 1024
        if len(spec.sizes_mb) == 1:
            target_filename = spec.filename
        else:
            target_filename = f"{stem}_{_size_label(int(size_mb))}{suffix}"
        path = out_dir / "docx" / target_filename
        budget = max(25_000, int(target_bytes * 0.985 / max(1, len(figs))))
        stats = {"words": 0, "figures": 0}
        for _attempt in range(3):
            encode_all(figs, budget, spec.prose.workers)
            stats = write_proposal_docx(
                path,
                spec,
                entities,
                prose,
                facts,
                figs,
                random.Random(seed + 1),
                header_filename=spec.filename,
                trap=False,
            )
            actual_size = path.stat().st_size
            if target_bytes <= actual_size <= target_bytes * 1.03:
                break
            budget = max(25_000, int(budget * (target_bytes / max(1, actual_size)) * 1.01))

        payload = path.read_bytes()
        sha256 = hashlib.sha256(payload).hexdigest()
        target_paths.append(path)
        manifest_entries.append(
            {
                "file": f"docx/{target_filename}",
                "size_mb": int(size_mb),
                "bytes": len(payload),
                "sha256": sha256,
                "words": stats["words"],
                "figures": stats["figures"],
                "trap_file": f"traps/{spec.trap_filename}",
                "trap_bytes": len(trap_bytes),
                "trap_sha256": trap_sha256,
                "seed": seed,
                "provider": spec.prose.provider,
                "model": spec.prose.model,
            }
        )

    primary_rel_file = str(manifest_entries[0]["file"])
    primary_payload = target_paths[0].read_bytes()

    eval_rows: list[dict[str, object]] = []
    for f in facts:
        locator = _resolve_text_locator(primary_payload, f.answer)
        eval_rows.append(
            {
                "id": f"{spec.id_prefix}-{f.fid}",
                "file": primary_rel_file,
                "modality": "text",
                "position": f.where,
                "question": f.question,
                "answer": f.answer,
                "trap_answer": f.trap_answer,
                "fact_id": f.fact_id,
                "trap_fact_id": f.trap_fact_id,
                "relation": f.relation,
                "evidence": f.sentence,
                "locator": locator.model_dump(),
            }
        )

    for (where, idx), (_, tf) in zip(picks.items(), trap_figs):
        fg = figs[idx]
        img_locator = DocxLocator(
            kind="image",
            unit_locator=f"image:{idx + 1}",
            index=idx + 1,
            caption=f"Plate {idx + 1}: {fg.caption}",
        )
        eval_rows.append(
            {
                "id": f"{spec.id_prefix}-I{idx + 1}",
                "file": primary_rel_file,
                "modality": "image",
                "position": where,
                "question": fg.question,
                "answer": fg.answer,
                "trap_answer": tf.answer,
                "fact_id": f"plate:{idx + 1}",
                "trap_fact_id": f"trap_plate:{idx + 1}",
                "relation": "plate_state",
                "evidence": f"image: Plate {idx + 1}: {fg.caption}",
                "locator": img_locator.model_dump(),
            }
        )

    with (out_dir / "manifest.jsonl").open("w", encoding="utf-8") as fh:
        for entry in manifest_entries:
            fh.write(json.dumps(entry) + "\n")

    with (out_dir / "evals.jsonl").open("w", encoding="utf-8") as fh:
        for row in eval_rows:
            fh.write(json.dumps(row) + "\n")

    with (out_dir / "evals.csv").open("w", encoding="utf-8", newline="") as fh:
        writer = csv.writer(fh)
        writer.writerow(["query", "golden"])
        for row in eval_rows:
            writer.writerow([row["question"], row["answer"]])

    return ProposalBuildResult(
        manifest_entries=manifest_entries,
        eval_rows=eval_rows,
        target_paths=target_paths,
        trap_path=trap_path,
    )
