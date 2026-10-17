"""Deck targets: WorldLoom enterprise deck, then a layout fix and an evidence slide.

The deck renderer assembles slides from several documents and ignores added IR sections.
So the text fact goes on a slide that BigLoom adds after rendering.
"""

from __future__ import annotations

from collections.abc import Sequence
from io import BytesIO
from typing import ClassVar

from ..evidence import pad_ooxml
from .base import Format, FormatHandler, Target


def pptx_add_evidence_slide(data: bytes, title: str, lines: Sequence[str], png: bytes | None) -> bytes:
    """Add one slide with text lines and an optional picture, at the middle of the deck."""
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation(BytesIO(data))
    layout = prs.slide_layouts[5] if len(prs.slide_layouts) > 5 else prs.slide_layouts[-1]
    slide = prs.slides.add_slide(layout)
    if slide.shapes.title is not None:
        slide.shapes.title.text = title
    box = slide.shapes.add_textbox(Inches(0.6), Inches(1.6), Inches(5.6), Inches(3.0))
    frame = box.text_frame
    frame.word_wrap = True
    for i, line in enumerate(lines):
        para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        para.text = line
        for run in para.runs:
            run.font.size = Pt(18)
    if png is not None:
        slide.shapes.add_picture(BytesIO(png), Inches(6.5), Inches(1.6), width=Inches(6.2))
    ids = prs.slides._sldIdLst
    new = ids[-1]
    ids.remove(new)
    ids.insert(len(ids) // 2, new)
    out = BytesIO()
    prs.save(out)
    return out.getvalue()


def fix_deck_layout(data: bytes) -> tuple[bytes, int]:
    """Give a valid size to every layout and slide placeholder with no height or wider than the slide.

    The WorldLoom enterprise deck template gives title and body placeholders a width of
    16 in and a height of 0 on a 13.33 in slide, so PowerPoint draws titles off the slide.
    Return the new bytes and the number of placeholders changed.
    """
    from pptx import Presentation
    from pptx.enum.shapes import PP_PLACEHOLDER
    from pptx.util import Inches

    prs = Presentation(BytesIO(data))
    w, h = prs.slide_width, prs.slide_height
    margin = Inches(0.6)
    inner = w - 2 * margin
    titles = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}
    changed = 0
    containers = [*prs.slide_layouts, *prs.slides]
    for container in containers:
        for ph in container.placeholders:
            if ph.height and ph.width and ph.left is not None and ph.left + ph.width <= w:
                continue
            kind = ph.placeholder_format.type
            if kind in titles:
                geometry = (margin, Inches(0.35), inner, Inches(1.0))
            elif kind == PP_PLACEHOLDER.DATE:
                geometry = (margin, h - Inches(0.55), Inches(3.0), Inches(0.35))
            elif kind == PP_PLACEHOLDER.FOOTER:
                geometry = ((w - Inches(6.0)) // 2, h - Inches(0.55), Inches(6.0), Inches(0.35))
            elif kind == PP_PLACEHOLDER.SLIDE_NUMBER:
                geometry = (w - margin - Inches(1.5), h - Inches(0.55), Inches(1.5), Inches(0.35))
            else:
                geometry = (margin, Inches(1.5), inner, h - Inches(2.2))
            ph.left, ph.top, ph.width, ph.height = geometry
            changed += 1
    out = BytesIO()
    prs.save(out)
    return out.getvalue(), changed


class PptxHandler(FormatHandler):
    fmt: ClassVar[Format] = "pptx"

    def add_picture(self, t: Target, payload: bytes, seed: int) -> bytes:
        payload, _ = fix_deck_layout(payload)
        return pptx_add_evidence_slide(payload, t.heading, [f"{t.heading}: {t.text_golden}"], t.chart(seed))

    def pad(self, payload: bytes, size_bytes: int, seed: int) -> bytes:
        return pad_ooxml(payload, size_bytes, seed)

    def finish_trap(self, t: Target, payload: bytes) -> bytes:
        payload, _ = fix_deck_layout(payload)
        lines = [f"{t.heading}: {t.text_canary}", f"{t.chart_title}: {t.image_canary}"]
        return pptx_add_evidence_slide(payload, f"{t.heading}{self.draft_suffix}", lines, None)
